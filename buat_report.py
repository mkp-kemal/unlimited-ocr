#!/usr/bin/env python
"""Buat report PDF biaya analisis risiko LLM dari log Excel (hasil/biaya-token-llm.xlsx).

Angkanya dihitung dari log, bukan diketik, supaya report bisa dibuat ulang setiap kali
uji dijalankan lagi dan selalu sama dengan data.

    .venv/bin/python buat_report.py
    .venv/bin/python buat_report.py --dokumen JASINDO --keluar hasil/report.pdf

Yang dilaporkan hanya analisis dengan tiga masukan — JSON hazard + scan OCR + system
prompt. Supaya yang dibandingkan hanya masukan yang identik, dipakai baris analisis dengan
system prompt, JSON hazard, dan scan OCR yang ukurannya sama dengan analisis TERBARU;
beberapa kali jalan dirata-rata.
"""

import argparse
import re
import unicodedata
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

from llm import MODEL_LLM

AKAR = Path(__file__).resolve().parent
NAMA = {i: n for i, n, *_ in MODEL_LLM}
UKURAN = ["karakter_system_prompt", "karakter_json_hazard", "karakter_scan_md"]

NAVY = colors.HexColor("#1F2A3C")
ABU = colors.HexColor("#F3F4F6")
GARIS = colors.HexColor("#D1D5DB")
HIJAU = colors.HexColor("#DCFCE7")


def baca_analisis(path, awalan):
    ws = load_workbook(path, read_only=True)["log"]
    baris = ws.iter_rows(values_only=True)
    df = pd.DataFrame(baris, columns=next(baris))
    a = df[(df["dokumen"].astype(str).str.startswith(awalan)) & (df["tahap"] == "analisis")]
    for k in UKURAN:
        a = a[pd.to_numeric(a[k], errors="coerce").fillna(0) > 0]
    if a.empty:
        raise SystemExit("Tidak ada analisis dengan tiga masukan (JSON hazard + scan OCR + "
                         "system prompt) untuk dokumen ini.")
    a = a.sort_values("waktu")
    akhir = a.iloc[-1]
    for k in UKURAN:
        a = a[a[k] == akhir[k]]
    return a, akhir


def baca_hasil_analisis(path):
    """Berkas hasil analisis (ditulis uji_biaya.py) -> {nama model: teks jawaban}."""
    hasil, nama, baris = {}, None, []
    for b in path.read_text(encoding="utf-8").splitlines():
        if b.startswith("### "):
            if nama:
                hasil[nama] = "\n".join(baris).strip().strip("-").strip()
            nama, baris = b[4:].strip(), []
        elif nama:
            baris.append(b)
    if nama:
        hasil[nama] = "\n".join(baris).strip().strip("-").strip()
    return hasil


def aman_font(teks):
    """Huruf bawaan PDF hanya punya karakter cp1252; di luar itu tercetak kotak hitam.

    Jawaban LLM tidak bisa diduga isinya, jadi karakter di luar cp1252 diganti padanan
    ASCII-nya (mis. panah) sebelum masuk PDF.
    """
    keluar = []
    for c in teks:
        try:
            c.encode("cp1252")
            keluar.append(c)
        except UnicodeEncodeError:
            keluar.append({"→": "->", "←": "<-", "≈": "~", "≤": "<=", "≥": ">=", "≠": "!="}.get(
                c, unicodedata.normalize("NFKD", c).encode("ascii", "ignore").decode() or "?"))
    return "".join(keluar)


def md_ke_paragraf(teks, gaya_teks, gaya_butir):
    """Markdown sederhana dari jawaban model (tebal, butir, kode) -> Paragraph reportlab."""
    hasil = []
    for b in aman_font(teks).splitlines():
        b = b.rstrip()
        if not b.strip():
            continue
        b = b.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        b = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", b)
        b = re.sub(r"`([^`]+)`", r"\1", b)
        # Miring satu bintang (*teks*), tapi bukan penanda butir "* " di awal baris.
        b = re.sub(r"(?<![*\w])\*(?!\s)(.+?)(?<!\s)\*(?![*\w])", r"<i>\1</i>", b)
        butir = re.match(r"^\s*[-*]\s+(.*)", b)
        if butir:
            hasil.append(Paragraph(butir.group(1), gaya_butir, bulletText="•"))
        else:
            hasil.append(Paragraph(b, gaya_teks))
    return hasil


def usd(x, desimal=6):
    return "—" if pd.isna(x) else f"${x:,.{desimal}f}"


def ribuan(x):
    return "—" if pd.isna(x) else f"{x:,.0f}".replace(",", ".")


def tabel(data, lebar, sorot=()):
    t = Table(data, colWidths=lebar, repeatRows=1)
    gaya = [
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, GARIS),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for r in range(2, len(data), 2):
        gaya.append(("BACKGROUND", (0, r), (-1, r), ABU))
    for r in sorot:
        gaya += [("BACKGROUND", (0, r), (-1, r), HIJAU),
                 ("FONTNAME", (0, r), (-1, r), "Helvetica-Bold")]
    t.setStyle(TableStyle(gaya))
    return t


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--log", type=Path, default=AKAR / "hasil" / "biaya-token-llm.xlsx")
    p.add_argument("--dokumen", default="JASINDO", help="awalan nama dokumen di log")
    p.add_argument("--keluar", type=Path, default=AKAR / "hasil" / "report-biaya-llm.pdf")
    p.add_argument("--analisis", type=Path,
                   help="berkas hasil analisis; bawaan hasil/analisis/<dokumen terbaru>.md")
    args = p.parse_args()

    a, akhir = baca_analisis(args.log, args.dokumen)
    g = a.groupby("model").agg(
        run=("biaya_usd", "size"), token_masuk=("token_masuk", "mean"),
        token_keluar=("token_keluar", "mean"), berpikir=("token_berpikir", "mean"),
        biaya=("biaya_usd", "mean"), biaya_min=("biaya_usd", "min"),
        biaya_max=("biaya_usd", "max"), detik=("detik_llm", "mean"),
    ).sort_values("biaya")

    termurah, tercepat, mahal = g["biaya"].idxmin(), g["detik"].idxmin(), g["biaya"].idxmax()
    halaman = int(akhir["halaman"]) if pd.notna(akhir["halaman"]) else None

    gaya = getSampleStyleSheet()
    judul = ParagraphStyle("judul", parent=gaya["Title"], fontSize=16, alignment=TA_LEFT,
                           spaceAfter=2)
    sub = ParagraphStyle("sub", parent=gaya["Normal"], fontSize=9,
                         textColor=colors.HexColor("#6B7280"))
    h2 = ParagraphStyle("h2", parent=gaya["Heading2"], fontSize=11.5, spaceBefore=10,
                        spaceAfter=5)
    teks = ParagraphStyle("teks", parent=gaya["Normal"], fontSize=9, leading=12.5)
    catatan = ParagraphStyle("catatan", parent=teks, fontSize=7.6, leading=10,
                             textColor=colors.HexColor("#4B5563"), fontName="Helvetica-Oblique")
    sel = ParagraphStyle("sel", parent=teks, fontSize=8.3, leading=10.5)

    run = sorted(set(g["run"].astype(int)))
    run_teks = f"{run[0]}" if len(run) == 1 else f"{run[0]}–{run[-1]}"
    cerita = [
        Paragraph("Ringkasan Biaya LLM: Analisis Risiko per Dokumen", judul),
        Paragraph(f"Dokumen {args.dokumen}{f' ({halaman} halaman)' if halaman else ''} · "
                  f"{len(g)} model · masukan: JSON hazard + scan OCR + system prompt · "
                  f"harga dan token dari OpenRouter · {date.today():%d-%m-%Y}", sub),
    ]

    # 1. Biaya per model
    data = [["Model", "Token\nmasuk", "Token\nkeluar", "— berpikir", "Biaya/\ndok (USD)",
             "Rentang\nbiaya/dok", "Per 1.000\ndok (USD)", "Waktu\n(s)"]]
    for m, r in g.iterrows():
        data.append([NAMA.get(m, m), ribuan(r.token_masuk), ribuan(r.token_keluar),
                     ribuan(r.berpikir), usd(r.biaya),
                     f"{usd(r.biaya_min, 5)}–{usd(r.biaya_max, 5)[1:]}",
                     usd(r.biaya * 1000, 2), f"{r.detik:.1f}"])
    cerita += [
        Paragraph("1. Biaya Analisis Risiko (per dokumen)", h2),
        tabel(data, [3.9 * cm, 1.6 * cm, 1.6 * cm, 1.6 * cm, 2.1 * cm, 2.6 * cm, 2.1 * cm, 1.4 * cm]),
        Spacer(1, 3),
        Paragraph(f"Rata-rata dari {run_teks} kali jalan per model dengan masukan identik. Token "
                  "\"berpikir\" sudah termasuk dalam token keluar dan tetap ditagih walau tidak "
                  "tampil di jawaban. OCR berjalan di komputer lokal dan tidak memakai token, "
                  "sehingga tidak termasuk biaya di atas.", catatan),
    ]

    # 2. Komposisi masukan
    ukuran = {"System prompt": int(akhir["karakter_system_prompt"]),
              "JSON hazard": int(akhir["karakter_json_hazard"]),
              f"Scan OCR{f' ({halaman} halaman)' if halaman else ''}": int(akhir["karakter_scan_md"])}
    total = sum(ukuran.values())
    data = [["Masukan", "Karakter", "Porsi"]]
    data += [[k, ribuan(v), f"{v / total:.0%}"] for k, v in ukuran.items()]
    data.append(["Total", ribuan(total), "100%"])
    porsi_scan = int(akhir["karakter_scan_md"]) / total
    cerita += [
        Paragraph("2. Komposisi Masukan (per dokumen)", h2),
        tabel(data, [6.0 * cm, 3.0 * cm, 3.0 * cm]),
        Spacer(1, 3),
        Paragraph(f"Scan OCR mengambil {porsi_scan:.0%} masukan, jadi biaya per dokumen terutama "
                  "naik-turun mengikuti panjang dokumen. System prompt sama persis di setiap "
                  "dokumen, sehingga bagian itu bisa di-cache penyedia model.", catatan),
    ]

    # 3. Kesimpulan
    data = [["Model", "Biaya/1.000 dok\n(USD)", "Waktu\n(s)", "Catatan"]]
    sorot = []
    for i, (m, r) in enumerate(g.iterrows(), 1):
        poin = []
        rasio = r.biaya / g.loc[termurah, "biaya"]
        if m == termurah:
            poin.append("termurah")
            sorot.append(i)
        elif rasio < 1.05:
            poin.append(f"biaya setara termurah (+{rasio - 1:.0%})")
        elif rasio < 2:
            poin.append(f"+{rasio - 1:.0%} dari termurah")
        else:
            poin.append(f"~{rasio:.0f}x biaya termurah")
        if m == tercepat:
            poin.append("tercepat")
        teks_poin = "; ".join(poin)
        data.append([NAMA.get(m, m), usd(r.biaya * 1000, 2), f"{r.detik:.1f}",
                     Paragraph(teks_poin[:1].upper() + teks_poin[1:], sel)])
    cerita += [Paragraph("Kesimpulan per Model", h2),
               tabel(data, [4.3 * cm, 3.2 * cm, 1.8 * cm, 7.6 * cm], sorot=sorot)]

    # Ringkasan dan batas
    cerita += [
        Spacer(1, 8),
        Paragraph(
            f"<b>Ringkasan:</b> analisis risiko satu dokumen dengan <b>{NAMA.get(termurah, termurah)}</b> "
            f"paling murah, sekitar <b>{usd(g.loc[termurah, 'biaya'] * 1000, 2)} per 1.000 dokumen</b>. "
            f"{NAMA.get(tercepat, tercepat)} paling cepat ({g.loc[tercepat, 'detik']:.1f} detik per "
            f"dokumen). {NAMA.get(mahal, mahal)} sekitar "
            f"{g.loc[mahal, 'biaya'] / g.loc[termurah, 'biaya']:.0f}x lebih mahal. Karena scan OCR "
            f"mengambil {porsi_scan:.0%} masukan, dokumen yang lebih panjang akan lebih mahal.", teks),
        Spacer(1, 8),
        Paragraph("<b>Batas report ini:</b> (1) satu dokumen, jadi belum mewakili dokumen dengan "
                  "panjang berbeda; (2) JSON hazard adalah contoh dari API dan koordinatnya tidak "
                  "sesuai alamat di dokumen; (3) system prompt analisis belum final — prompt yang "
                  "lebih panjang menaikkan token masuk; (4) biaya dalam USD, diambil dari respons "
                  "OpenRouter; (5) kualitas isi analisis belum dinilai, report ini hanya mengukur "
                  "biaya.", catatan),
    ]

    # Hasil analisis, di bawah bagian biaya
    path_analisis = args.analisis or AKAR / "hasil" / "analisis" / f"{Path(str(akhir['dokumen'])).stem}.md"
    if path_analisis.exists():
        butir = ParagraphStyle("butir", parent=teks, leftIndent=12, bulletIndent=2, spaceBefore=1)
        isi_teks = ParagraphStyle("isi", parent=teks, spaceBefore=4)
        model_h = ParagraphStyle("model_h", parent=h2, fontSize=10.5, spaceBefore=12,
                                 textColor=NAVY)
        cerita += [PageBreak(), Paragraph("Hasil Analisis per Model", h2),
                   Paragraph("Jawaban dari kali jalan terakhir, ditampilkan apa adanya. Isi bisa "
                             "sedikit berbeda antar jalan.", catatan)]
        jawaban = baca_hasil_analisis(path_analisis)
        for m in g.index:
            nama = NAMA.get(m, m)
            if nama in jawaban:
                isi = md_ke_paragraf(jawaban[nama], isi_teks, butir)
                # Judul model tidak boleh tertinggal sendirian di dasar halaman.
                cerita += [KeepTogether([Paragraph(nama, model_h)] + isi[:2])] + isi[2:]
    else:
        print(f"Catatan: {path_analisis} tidak ada — bagian hasil analisis dilewati.")

    args.keluar.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(args.keluar), pagesize=A4, leftMargin=1.7 * cm, rightMargin=1.7 * cm,
                      topMargin=1.4 * cm, bottomMargin=1.4 * cm,
                      title="Ringkasan Biaya LLM — Analisis Risiko",
                      author="unlimited-ocr").build(cerita)
    print(f"Report: {args.keluar}")


if __name__ == "__main__":
    main()

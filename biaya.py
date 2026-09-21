"""Catat token dan biaya LLM ke satu file Excel yang terus bertambah.

Satu baris = satu panggilan LLM: tahap `ekstraksi` (ambil field dari teks OCR) atau
`analisis` (analisis risiko dari JSON hazard + scan OCR). Sheet `ringkasan` menjumlahkan keduanya
menjadi biaya per dokumen per model, dengan rumus — bukan angka tertulis — supaya tetap
benar kalau baris log disunting atau dihapus di Excel.
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from llm import MODEL_LLM

FILE = Path(__file__).resolve().parent / "hasil" / "biaya-token-llm.xlsx"

# (judul kolom, lebar, format angka)
KOLOM = [
    ("waktu", 20, "yyyy-mm-dd hh:mm:ss"),
    ("tahap", 11, None),
    ("dokumen", 28, None),
    ("mesin_ocr", 22, None),
    ("halaman", 9, "0"),
    # Besar teks yang dikirim menentukan token masuk; tanpa kolom ini, beda biaya antar
    # dokumen tidak bisa dijelaskan.
    ("karakter_dikirim", 16, "#,##0"),
    # System prompt dikirim ulang di setiap dokumen, jadi porsinya dicatat terpisah.
    ("karakter_system_prompt", 22, "#,##0"),
    ("model", 30, None),
    ("tingkat", 10, None),
    ("field_diminta", 13, "0"),
    ("field_ditemukan", 15, "0"),
    ("field_terbukti", 14, "0"),
    ("token_masuk", 13, "#,##0"),
    ("token_cache", 12, "#,##0"),
    ("token_keluar", 13, "#,##0"),
    ("token_berpikir", 15, "#,##0"),
    ("total_token", 12, "#,##0"),
    ("biaya_usd", 13, "$0.000000"),
    ("biaya_per_1000_dokumen_usd", 26, "$#,##0.00"),
    ("detik_llm", 10, "0.00"),
    ("status", 22, None),
    # Rincian ukuran masukan analisis. Ditaruh di ujung, bukan di samping karakter lain,
    # supaya baris lama di file Excel yang sudah ada tidak bergeser kolom.
    ("karakter_json_hazard", 20, "#,##0"),
    ("karakter_scan_md", 17, "#,##0"),
]
HURUF = {nama: get_column_letter(i + 1) for i, (nama, _, _) in enumerate(KOLOM)}

TEBAL = Font(bold=True, color="FFFFFF")
LATAR = PatternFill("solid", fgColor="1F4E78")


def _kepala(ws, judul):
    for i, (nama, lebar, _) in enumerate(judul, 1):
        sel = ws.cell(row=1, column=i, value=nama)
        sel.font, sel.fill = TEBAL, LATAR
        sel.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = lebar
    ws.freeze_panes = "A2"


def _buka():
    if FILE.exists():
        return load_workbook(FILE)
    FILE.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "log"
    _kepala(ws, KOLOM)
    return wb


def _ringkasan(wb):
    """Sheet ringkasan dibangun ulang tiap kali, supaya daftar modelnya selalu terkini."""
    if "ringkasan" in wb.sheetnames:
        del wb["ringkasan"]
    ws = wb.create_sheet("ringkasan")
    judul = [
        ("model", 30, None),
        ("tingkat", 10, None),
        ("run_ekstraksi", 14, "0"),
        ("rata2_biaya_ekstraksi_usd", 24, "$0.000000"),
        ("run_analisis", 13, "0"),
        ("rata2_biaya_analisis_usd", 23, "$0.000000"),
        ("rata2_token_masuk_analisis", 25, "#,##0"),
        ("total_per_dokumen_usd", 21, "$0.000000"),
        ("per_1000_dokumen_usd", 20, "$#,##0.00"),
    ]
    _kepala(ws, judul)

    def kol(k):
        return f"log!${HURUF[k]}:${HURUF[k]}"

    m, t = kol("model"), kol("tahap")
    for r, (id_model, _, tingkat, *_) in enumerate(MODEL_LLM, 2):
        ws.cell(row=r, column=1, value=id_model)
        ws.cell(row=r, column=2, value=tingkat)
        ws.cell(row=r, column=3, value=f'=COUNTIFS({m},A{r},{t},"ekstraksi")')
        # IFERROR: model yang belum pernah dijalankan tidak punya rata-rata, dan #DIV/0!
        # di lembar ringkasan lebih membingungkan daripada sel kosong.
        ws.cell(row=r, column=4,
                value=f'=IFERROR(AVERAGEIFS({kol("biaya_usd")},{m},A{r},{t},"ekstraksi"),"")')
        ws.cell(row=r, column=5, value=f'=COUNTIFS({m},A{r},{t},"analisis")')
        ws.cell(row=r, column=6,
                value=f'=IFERROR(AVERAGEIFS({kol("biaya_usd")},{m},A{r},{t},"analisis"),"")')
        ws.cell(row=r, column=7,
                value=f'=IFERROR(AVERAGEIFS({kol("token_masuk")},{m},A{r},{t},"analisis"),"")')
        # Total hanya diisi kalau model itu sudah menjalani KEDUA tahap. Menjumlahkan
        # satu tahap saja akan tampak seperti biaya per dokumen, padahal separuhnya.
        ws.cell(row=r, column=8, value=f'=IF(AND(C{r}>0,E{r}>0),D{r}+F{r},"")')
        ws.cell(row=r, column=9, value=f'=IFERROR(H{r}*1000,"")')
        for c, (_, _, fmt) in enumerate(judul, 1):
            if fmt:
                ws.cell(row=r, column=c).number_format = fmt


def catat(tahap, dokumen, mesin_ocr, halaman, karakter_dikirim, model, pemakaian,
          karakter_system_prompt=0, karakter_json_hazard=None, karakter_scan_md=None,
          field_diminta=None, field_ditemukan=None, field_terbukti=None, status="ok"):
    """Tambahkan satu baris ke log. Mengembalikan path file Excel."""
    tingkat = next((tk for i, _, tk, *_ in MODEL_LLM if i == model), "")
    wb = _buka()
    ws = wb["log"]
    r = ws.max_row + 1

    nilai = {
        "waktu": datetime.now().replace(microsecond=0),
        "tahap": tahap,
        "dokumen": dokumen,
        "mesin_ocr": mesin_ocr,
        "halaman": halaman,
        "karakter_dikirim": karakter_dikirim,
        "karakter_system_prompt": karakter_system_prompt,
        "model": model,
        "tingkat": tingkat,
        "field_diminta": field_diminta,
        "field_ditemukan": field_ditemukan,
        "field_terbukti": field_terbukti,
        "token_masuk": pemakaian.get("token_masuk", 0),
        "token_cache": pemakaian.get("token_cache", 0),
        "token_keluar": pemakaian.get("token_keluar", 0),
        "token_berpikir": pemakaian.get("token_berpikir", 0),
        "total_token": pemakaian.get("total_token", 0),
        "biaya_usd": pemakaian.get("biaya_usd", 0.0),
        # Rumus, bukan angka: tetap benar kalau biaya_usd dikoreksi manual di Excel.
        "biaya_per_1000_dokumen_usd": f"={HURUF['biaya_usd']}{r}*1000",
        "detik_llm": pemakaian.get("detik", 0.0),
        "status": status,
        "karakter_json_hazard": karakter_json_hazard,
        "karakter_scan_md": karakter_scan_md,
    }
    for c, (nama, _, fmt) in enumerate(KOLOM, 1):
        sel = ws.cell(row=r, column=c, value=nilai[nama])
        if fmt:
            sel.number_format = fmt

    _ringkasan(wb)
    wb.save(FILE)
    return FILE
